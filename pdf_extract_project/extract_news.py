import argparse
import json
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import fitz
from docx import Document
from docx.shared import Pt

TOP_MARGIN = 55
BOTTOM_MARGIN = 715
INDENT_THRESHOLD = 7.5

BYLINE_RE = re.compile(r"^(By|Interviewed by|Written by)\b", re.IGNORECASE)
CAPTION_MARKERS = (
    "photo by",
    "photos by",
    "illustration by",
    "image by",
    "source:",
    "sources:",
)
TEXT_FIXES = {
    "ﬁ": "fi",
    "ﬂ": "fl",
    "’": "'",
    "“": '"',
    "”": '"',
    "–": "-",
    "—": "-",
    "…": "...",
    "\u00a0": " ",
}


@dataclass
class LineInfo:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    avg_size: float
    max_size: float


@dataclass
class BlockInfo:
    page_index: int
    block_no: int
    x0: float
    y0: float
    x1: float
    y1: float
    lines: list[LineInfo]
    raw_text: str
    avg_size: float
    max_size: float
    prefix: str = ""

    @property
    def text(self) -> str:
        return self.raw_text

    @property
    def width(self) -> float:
        return self.x1 - self.x0

    @property
    def height(self) -> float:
        return self.y1 - self.y0

    @property
    def word_count(self) -> int:
        return len(self.text.replace("\n", " ").split())


@dataclass
class ArticleSeed:
    start_page_index: int
    section: str
    title: str
    subtitle: str
    byline: str
    title_block_no: int
    subtitle_block_no: int | None
    byline_block_no: int | None
    body_start_y: float
    kind: str = "feature"
    title_x: float = 0.0
    title_y: float = 0.0


@dataclass
class Article:
    index: int
    section: str
    title: str
    subtitle: str
    byline: str
    start_pdf_page: int
    end_pdf_page: int
    paragraphs: list[str] = field(default_factory=list)
    sort_page_index: int = 0
    sort_y: float = 0.0
    sort_x: float = 0.0

    @property
    def page_range(self) -> str:
        return str(self.start_pdf_page) if self.start_pdf_page == self.end_pdf_page else f"{self.start_pdf_page}-{self.end_pdf_page}"


@dataclass
class OrderedLine:
    text: str
    x0: float
    y0: float
    y1: float
    column: int
    baseline_x: float
    is_heading: bool = False


@dataclass
class ShortArticleCandidate:
    section: str
    title: str
    subtitle: str
    page_index: int
    title_y: float
    blocks: list[BlockInfo]
    layout: str = "column"
    title_x: float = 0.0
    paragraphs_override: list[str] = field(default_factory=list)


def normalize_text(text: str) -> str:
    for src, dst in TEXT_FIXES.items():
        text = text.replace(src, dst)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_line(text: str) -> str:
    return normalize_text(text).replace("|", "").strip()


def extract_blocks(page: fitz.Page) -> list[BlockInfo]:
    data = page.get_text("dict")
    blocks: list[BlockInfo] = []
    for block_no, block in enumerate(data.get("blocks", [])):
        if block.get("type") != 0:
            continue
        lines: list[LineInfo] = []
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = normalize_line("".join(span.get("text", "") for span in spans))
            if not text:
                continue
            sizes = [float(span.get("size", 0.0)) for span in spans] or [0.0]
            bbox = line["bbox"]
            lines.append(LineInfo(text, float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3]), sum(sizes) / len(sizes), max(sizes)))
        if not lines:
            continue
        bbox = block["bbox"]
        blocks.append(BlockInfo(page.number, block_no, float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3]), lines, "\n".join(line.text for line in lines), sum(line.avg_size for line in lines) / len(lines), max(line.max_size for line in lines)))
    return blocks


def is_header_or_footer(block: BlockInfo) -> bool:
    return block.y0 < TOP_MARGIN or block.y1 > BOTTOM_MARGIN


def is_section_label(block: BlockInfo) -> bool:
    text = block.text.strip()
    if not text or block.y0 > 85 or block.word_count > 6:
        return False
    letters = [ch for ch in text if ch.isalpha()]
    if not letters:
        return False
    ratio = sum(1 for ch in letters if ch.isupper()) / len(letters)
    if ratio >= 0.75:
        return True
    return block.y0 <= 30 and block.max_size >= 20 and block.word_count <= 4


def page_section_label(blocks: list[BlockInfo]) -> str:
    labels = [block.text.strip() for block in blocks if is_section_label(block)]
    return labels[0] if labels else ""


def is_byline(block: BlockInfo) -> bool:
    text = block.text.strip()
    return block.word_count <= 8 and not any(ch.isdigit() for ch in text) and bool(BYLINE_RE.match(text))


def is_drop_cap(block: BlockInfo) -> bool:
    text = block.text.strip()
    return len(text) == 1 and text.isalpha() and block.max_size >= 30 and block.width < 60


def is_decorative_block(block: BlockInfo) -> bool:
    punctuation = "".join(ch for ch in block.text.strip() if not ch.isspace())
    return bool(punctuation) and len(punctuation) <= 3 and all(not ch.isalnum() for ch in punctuation) and block.max_size >= 16


def is_caption(block: BlockInfo, page_width: float) -> bool:
    text = block.text.lower().strip()
    if not text or is_byline(block):
        return False
    if any(marker in text for marker in CAPTION_MARKERS):
        return True
    if block.max_size <= 8.6 and block.height < 44 and block.word_count <= 100:
        return True
    return block.max_size <= 8.1 and block.word_count <= 140 and block.width < page_width * 0.95


def is_probable_title_block(block: BlockInfo) -> bool:
    return block.y0 < 620 and block.max_size >= 14 and block.word_count <= 22 and not is_section_label(block)


def format_title(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return ""
    if len(lines) == 1:
        return lines[0]
    if any(line and line[0].islower() for line in lines[1:]):
        return " ".join(lines)
    return ": ".join(lines)


def split_block_by_x_gaps(block: BlockInfo) -> list[BlockInfo]:
    if len(block.lines) <= 1:
        return [block]
    centers: list[float] = []
    groups: dict[int, list[LineInfo]] = {}
    for line in block.lines:
        center = (line.x0 + line.x1) / 2
        idx = None
        for i, existing in enumerate(centers):
            if abs(center - existing) <= 70:
                idx = i
                centers[i] = (centers[i] + center) / 2
                break
        if idx is None:
            idx = len(centers)
            centers.append(center)
        groups.setdefault(idx, []).append(line)
    if len(groups) <= 1:
        return [block]
    parts: list[BlockInfo] = []
    for idx in sorted(groups):
        lines = groups[idx]
        parts.append(BlockInfo(block.page_index, block.block_no * 10 + idx, min(line.x0 for line in lines), min(line.y0 for line in lines), max(line.x1 for line in lines), max(line.y1 for line in lines), lines, "\n".join(line.text for line in lines), sum(line.avg_size for line in lines) / len(lines), max(line.max_size for line in lines)))
    return parts


def month_names_present(text: str) -> bool:
    return bool(re.search(r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\b", text, re.IGNORECASE))


def title_case_score(text: str) -> float:
    words = re.findall(r"[A-Za-z][A-Za-z'’\-]*", text)
    if not words:
        return 0.0
    ignored = {"a", "an", "the", "and", "or", "of", "to", "for", "in", "on", "at", "by", "from", "with", "into", "over", "under", "after"}
    scored = [word for word in words if word.lower() not in ignored] or words
    return sum(1 for word in scored if word[0].isupper() or word.isupper()) / len(scored)


def looks_like_short_title(block: BlockInfo) -> bool:
    text = block.text.strip()
    return bool(text) and block.word_count <= 24 and block.height <= 120 and title_case_score(text) >= 0.55 and not (text.count('"') >= 2 and block.word_count > 12)


def line_looks_like_meta_text(text: str) -> bool:
    lower = text.lower()
    return month_names_present(text) or any(marker in lower for marker in ("weekly", "magazine", "edition", "outlook", "newschina", "newsweek", "world vision", "guan cha"))


def split_short_title_parts(block: BlockInfo) -> tuple[str, str]:
    lines = [normalize_text(line.text) for line in block.lines if normalize_text(line.text)]
    if not lines:
        return "", ""
    if len(lines) == 1:
        return format_title(lines[0]), ""
    meta_lines: list[str] = []
    title_lines = lines[:]
    date_idx = next((idx for idx, line in enumerate(lines[:3]) if month_names_present(line)), None)
    if date_idx is not None and date_idx + 1 < len(lines):
        meta_lines = lines[: date_idx + 1]
        title_lines = lines[date_idx + 1 :]
    while len(title_lines) > 1 and line_looks_like_meta_text(title_lines[0]):
        meta_lines.append(title_lines.pop(0))
    return format_title("\n".join(title_lines)), " | ".join(dict.fromkeys(meta_lines))


def is_short_news_source_block(block: BlockInfo) -> bool:
    text = block.text.strip()
    return bool(text) and block.max_size <= 12 and block.word_count <= 10 and month_names_present(text)


def is_short_meta_block(block: BlockInfo) -> bool:
    text = block.text.strip()
    if not text or block.max_size > 10.5 or block.height > 28 or block.word_count > 10:
        return False
    if month_names_present(text) or any(ch.isdigit() for ch in text):
        return True
    words = re.findall(r"[A-Za-z][A-Za-z'’\-]*", text)
    return 0 < len(words) <= 4 and all(word[0].isupper() for word in words)


def is_short_title_block(block: BlockInfo, page_width: float, section: str) -> bool:
    if is_section_label(block) or is_byline(block) or is_decorative_block(block) or is_short_news_source_block(block):
        return False
    if block.y0 < 32 or block.y1 > BOTTOM_MARGIN:
        return False
    if is_caption(block, page_width) and section != "NETIZENWATCH":
        return False
    if section == "MEDIAFOCUS":
        return looks_like_short_title(block) and block.max_size >= 9.3 and block.height <= 42 and block.word_count <= 16
    return looks_like_short_title(block) and (block.max_size >= 12 or (block.max_size >= 10.2 and block.height <= 70 and block.page_index + 1 <= 15))


def candidate_range(title_block: BlockInfo, page_width: float) -> tuple[float, float]:
    return max(0.0, title_block.x0 - 12), min(page_width, title_block.x1 + 16)


def split_short_candidate_blocks(block: BlockInfo) -> list[BlockInfo]:
    return split_block_by_x_gaps(block) if block.max_size >= 11 or block.height <= 90 else [block]


def median_body_size(blocks: Iterable[BlockInfo]) -> float:
    sizes = [line.avg_size for block in blocks for line in block.lines if 8.5 <= line.avg_size <= 11.5]
    return statistics.median(sizes) if sizes else 10.0


def detect_article_seed(page: fitz.Page, blocks: list[BlockInfo]) -> ArticleSeed | None:
    section = page_section_label(blocks)
    if page.number <= 3 and section == "EDITORIAL":
        editorial_blocks = [block for block in blocks if not is_decorative_block(block) and not is_caption(block, page.rect.width)]
        title_candidates = [block for block in editorial_blocks if is_probable_title_block(block) and not is_drop_cap(block) and not is_section_label(block)]
        if title_candidates:
            title_block = max(title_candidates, key=lambda block: (block.max_size, block.width))
            return ArticleSeed(page.number, "EDITORIAL", format_title(title_block.text), "", "", title_block.block_no, None, None, title_block.y1 + 8, "editorial", title_block.x0, title_block.y0)

    content_blocks = [block for block in blocks if not is_header_or_footer(block)]
    bylines = [block for block in content_blocks if is_byline(block) and 80 <= block.y0 <= 680]
    if bylines:
        byline = min(bylines, key=lambda item: item.y0)
        title_candidates = [block for block in content_blocks if block.y0 < byline.y0 and is_probable_title_block(block) and not is_caption(block, page.rect.width) and not is_drop_cap(block) and not is_byline(block)]
        if not title_candidates:
            return None
        title_block = max(title_candidates, key=lambda block: (block.max_size, block.width))
        subtitle_candidates = [block for block in content_blocks if block.block_no not in {title_block.block_no, byline.block_no} and title_block.y1 <= block.y0 <= byline.y0 and block.max_size >= 10 and block.word_count <= 45 and not is_caption(block, page.rect.width) and not is_section_label(block)]
        subtitle_block = min(subtitle_candidates, key=lambda item: item.y0) if subtitle_candidates else None
        return ArticleSeed(page.number, section, format_title(title_block.text), normalize_text(subtitle_block.text) if subtitle_block else "", normalize_text(byline.text), title_block.block_no, subtitle_block.block_no if subtitle_block else None, byline.block_no, byline.y1 + 12, "feature", title_block.x0, title_block.y0)

    if page.number + 1 >= 16:
        title_candidates = [block for block in content_blocks if is_probable_title_block(block) and not is_drop_cap(block) and not is_caption(block, page.rect.width) and not is_byline(block)]
        if title_candidates:
            title_block = max(title_candidates, key=lambda block: (block.max_size, block.width))
            subtitle_candidates = [block for block in content_blocks if block.block_no != title_block.block_no and title_block.y1 <= block.y0 <= title_block.y1 + 180 and block.max_size >= 11 and block.word_count <= 50 and not is_caption(block, page.rect.width) and not is_section_label(block)]
            subtitle_block = min(subtitle_candidates, key=lambda item: item.y0) if subtitle_candidates else None
            byline_candidates = [block for block in content_blocks if is_byline(block) and title_block.y1 <= block.y0 <= title_block.y1 + 220]
            byline_block = min(byline_candidates, key=lambda item: item.y0) if byline_candidates else None
            if title_block.max_size >= 32 and subtitle_block is not None:
                body_start_y = (byline_block.y1 + 12) if byline_block else (subtitle_block.y1 + 16)
                return ArticleSeed(page.number, section, format_title(title_block.text), normalize_text(subtitle_block.text), normalize_text(byline_block.text) if byline_block else "", title_block.block_no, subtitle_block.block_no, byline_block.block_no if byline_block else None, body_start_y, "feature", title_block.x0, title_block.y0)
    return None

def assign_drop_caps(blocks: list[BlockInfo]) -> list[BlockInfo]:
    content = [block for block in blocks if not is_drop_cap(block)]
    drops = [block for block in blocks if is_drop_cap(block)]
    for drop in drops:
        candidates = [
            block
            for block in content
            if block.x0 >= drop.x0 - 12
            and abs(block.y0 - drop.y0) <= 40
            and block.x0 - drop.x1 <= 140
        ]
        if candidates:
            target = min(candidates, key=lambda block: (abs(block.y0 - drop.y0), block.x0))
            if not target.prefix:
                target.prefix = drop.text
    return content


def build_page_body_blocks(page: fitz.Page, blocks: list[BlockInfo], seed: ArticleSeed, is_first_page: bool) -> list[BlockInfo]:
    filtered: list[BlockInfo] = []
    for block in blocks:
        if is_header_or_footer(block) or is_decorative_block(block) or is_caption(block, page.rect.width) or is_section_label(block):
            continue
        if seed.kind == "editorial" and block.x0 >= page.rect.width * 0.68:
            continue
        if is_first_page and block.block_no in {seed.title_block_no, seed.subtitle_block_no, seed.byline_block_no}:
            continue
        if is_first_page and block.y0 < seed.body_start_y and not is_drop_cap(block):
            continue
        if not is_first_page and block.y0 < TOP_MARGIN + 5:
            continue
        filtered.append(block)
    return assign_drop_caps(filtered)


def cluster_columns(blocks: list[BlockInfo]) -> list[float]:
    xs = sorted(block.x0 for block in blocks)
    if not xs:
        return []
    columns = [xs[0]]
    for x in xs[1:]:
        if abs(x - columns[-1]) > 70:
            columns.append(x)
        else:
            columns[-1] = (columns[-1] + x) / 2
    return columns


def order_blocks(page: fitz.Page, blocks: list[BlockInfo]) -> list[BlockInfo]:
    usable_width = page.rect.width - 50
    narrow = [block for block in blocks if block.width < usable_width * 0.78]
    full = [block for block in blocks if block.width >= usable_width * 0.78]
    if not narrow:
        return sorted(full, key=lambda block: (block.y0, block.x0))
    columns = cluster_columns(narrow)

    def column_index(block: BlockInfo) -> int:
        return min(range(len(columns)), key=lambda idx: abs(block.x0 - columns[idx])) if columns else 0

    top_of_columns = min(block.y0 for block in narrow)
    leading_full = [block for block in full if block.y1 <= top_of_columns + 8]
    trailing_full = [block for block in full if block not in leading_full]
    ordered = sorted(leading_full, key=lambda block: (block.y0, block.x0))
    ordered.extend(sorted(narrow, key=lambda block: (column_index(block), block.y0)))
    ordered.extend(sorted(trailing_full, key=lambda block: (block.y0, block.x0)))
    return ordered


def line_is_heading(text: str, line: LineInfo, body_size: float) -> bool:
    return line.max_size >= body_size + 1 and len(text.split()) <= 8 and not text.endswith((".", ",", ";", ":", "?", "!"))


def build_ordered_lines(page: fitz.Page, blocks: list[BlockInfo], body_size: float) -> list[OrderedLine]:
    ordered_blocks = order_blocks(page, blocks)
    usable_width = page.rect.width - 50
    narrow = [block for block in blocks if block.width < usable_width * 0.78]
    columns = cluster_columns(narrow) or [min((block.x0 for block in ordered_blocks), default=0.0)]

    def column_index(block: BlockInfo) -> int:
        return min(range(len(columns)), key=lambda idx: abs(block.x0 - columns[idx]))

    baseline_by_column: dict[int, float] = {}
    for block in ordered_blocks:
        idx = column_index(block)
        baseline_by_column[idx] = min(baseline_by_column.get(idx, block.x0), block.x0)

    lines: list[OrderedLine] = []
    for block in ordered_blocks:
        idx = column_index(block)
        baseline_x = baseline_by_column.get(idx, block.x0)
        for line_index, line in enumerate(block.lines):
            text = line.text
            if line_index == 0 and block.prefix:
                text = block.prefix + text
            text = normalize_text(text)
            if text:
                lines.append(OrderedLine(text, line.x0, line.y0, line.y1, idx, baseline_x, line_is_heading(text, line, body_size)))
    return lines


def build_single_column_ordered_lines(blocks: list[BlockInfo], body_size: float) -> list[OrderedLine]:
    lines: list[OrderedLine] = []
    for block in sorted(blocks, key=lambda item: (item.y0, item.x0)):
        baseline_x = min((line.x0 for line in block.lines), default=block.x0)
        for line_index, line in enumerate(block.lines):
            text = line.text
            if line_index == 0 and block.prefix:
                text = block.prefix + text
            text = normalize_text(text)
            if text:
                lines.append(OrderedLine(text, line.x0, line.y0, line.y1, 0, baseline_x, line_is_heading(text, line, body_size)))
    return lines


def append_line(current: str, text: str) -> str:
    return current[:-1] + text if current.endswith("-") and text and text[0].islower() else f"{current} {text}"


def join_lines(lines: Iterable[str]) -> str:
    current = ""
    for line in lines:
        text = normalize_text(line)
        if not text:
            continue
        current = text if not current else append_line(current, text)
    return normalize_text(current)


def split_verbatim_block(block: BlockInfo) -> tuple[str, str]:
    lines = [normalize_text(line.text) for line in block.lines if normalize_text(line.text)]
    if not lines:
        return "", ""

    quote_lines: list[str] = []
    attribution_lines: list[str] = []
    quote_closed = False
    for line in lines:
        if quote_closed:
            attribution_lines.append(line)
            continue
        quote_lines.append(line)
        if "”" in line or line.endswith('"'):
            quote_closed = True

    if not attribution_lines and len(lines) >= 2:
        quote_lines = lines[:-1]
        attribution_lines = [lines[-1]]

    quote_text = join_lines(quote_lines)
    if not attribution_lines:
        return quote_text, ""

    attribution_text = join_lines(attribution_lines)
    return quote_text, attribution_text


def lines_to_paragraphs(lines: list[OrderedLine]) -> list[str]:
    paragraphs: list[str] = []
    current = ""
    previous_line: OrderedLine | None = None
    for line in lines:
        if line.is_heading:
            if current:
                paragraphs.append(normalize_text(current))
                current = ""
            paragraphs.append(normalize_text(line.text))
            previous_line = None
            continue
        if not current:
            current = line.text
            previous_line = line
            continue
        current_indent = line.x0 - line.baseline_x
        previous_indent = 0.0 if previous_line is None else previous_line.x0 - previous_line.baseline_x
        repeated_indent = previous_line is not None and current_indent >= INDENT_THRESHOLD and previous_indent >= INDENT_THRESHOLD and abs(current_indent - previous_indent) < 3
        if current_indent >= INDENT_THRESHOLD and not repeated_indent:
            paragraphs.append(normalize_text(current))
            current = line.text
        else:
            current = append_line(current, line.text)
        previous_line = line
    if current:
        paragraphs.append(normalize_text(current))
    return [paragraph for paragraph in paragraphs if paragraph]


def build_article(doc: fitz.Document, seed: ArticleSeed, end_page_index: int, index: int) -> Article:
    page_blocks_cache: list[list[BlockInfo]] = []
    body_sizes: list[float] = []
    for page_index in range(seed.start_page_index, end_page_index + 1):
        page = doc[page_index]
        blocks = extract_blocks(page)
        filtered = build_page_body_blocks(page, blocks, seed, is_first_page=page_index == seed.start_page_index)
        page_blocks_cache.append(filtered)
        body_sizes.append(median_body_size(filtered))
    body_size = statistics.median(body_sizes) if body_sizes else 10.0
    ordered_lines: list[OrderedLine] = []
    for offset, page_index in enumerate(range(seed.start_page_index, end_page_index + 1)):
        ordered_lines.extend(build_ordered_lines(doc[page_index], page_blocks_cache[offset], body_size))
    return Article(index, seed.section, seed.title, seed.subtitle, seed.byline, seed.start_page_index + 1, end_page_index + 1, lines_to_paragraphs(ordered_lines), seed.start_page_index, seed.title_y, seed.title_x)


def build_short_article(candidate: ShortArticleCandidate, index: int) -> Article:
    if candidate.paragraphs_override:
        return Article(index, candidate.section, candidate.title, candidate.subtitle, "", candidate.page_index + 1, candidate.page_index + 1, candidate.paragraphs_override, candidate.page_index, candidate.title_y, candidate.title_x)
    page_width = max((block.x1 for block in candidate.blocks), default=540.0) + 30
    samples = [line.avg_size for block in candidate.blocks for line in block.lines if 8.0 <= line.avg_size <= 11.5]
    body_size = statistics.median(samples) if samples else 10.0
    pseudo_page = type("PseudoPage", (), {"rect": fitz.Rect(0.0, 0.0, page_width, 750.0)})()
    ordered_lines = build_single_column_ordered_lines(candidate.blocks, body_size) if candidate.layout == "single" else build_ordered_lines(pseudo_page, candidate.blocks, body_size)
    sort_y = candidate.title_x if candidate.section == "MEDIAFOCUS" else candidate.title_y
    sort_x = candidate.title_y if candidate.section == "MEDIAFOCUS" else candidate.title_x
    return Article(index, candidate.section, candidate.title, candidate.subtitle, "", candidate.page_index + 1, candidate.page_index + 1, lines_to_paragraphs(ordered_lines), candidate.page_index, sort_y, sort_x)


def detect_break_pages(doc: fitz.Document) -> set[int]:
    break_pages: set[int] = set()
    for page_index in range(doc.page_count):
        page = doc[page_index]
        blocks = extract_blocks(page)
        labels = {block.text.strip().upper() for block in blocks}
        if "CONTENTS" in labels or "SNAPSHOTS" in labels:
            break_pages.add(page_index)
            continue
        section_blocks = [block for block in blocks if is_section_label(block)]
        if not section_blocks:
            continue
        visible = [block for block in blocks if not is_header_or_footer(block) and not is_decorative_block(block) and not is_caption(block, page.rect.width)]
        if any(is_byline(block) for block in visible):
            continue
        body_blocks = [block for block in visible if not is_section_label(block)]
        body_words = sum(block.word_count for block in body_blocks)
        first_body_y = min((block.y0 for block in body_blocks), default=999)
        if not body_blocks or body_words < 60 or (first_body_y > 150 and len(body_blocks) <= 3):
            break_pages.add(page_index)
    return break_pages


def extract_verbatim_blocks(blocks: list[BlockInfo]) -> list[BlockInfo]:
    candidates = [
        block
        for block in blocks
        if not is_header_or_footer(block)
        and not is_section_label(block)
        and not is_decorative_block(block)
        and block.word_count >= 18
        and block.text.strip().startswith(("“", '"'))
    ]
    if not candidates:
        return []
    columns = cluster_columns(candidates)

    def column_index(block: BlockInfo) -> int:
        return min(range(len(columns)), key=lambda idx: abs(block.x0 - columns[idx])) if columns else 0

    return sorted(candidates, key=lambda block: (column_index(block), block.y0))


def build_verbatim_article(page_blocks_cache: dict[int, list[BlockInfo]]) -> Article | None:
    verbatim_pages = [page_index for page_index, blocks in page_blocks_cache.items() if page_section_label(blocks) == "VERBATIM"]
    if not verbatim_pages:
        return None

    paragraphs: list[str] = []
    first_page = min(verbatim_pages)
    last_page = max(verbatim_pages)
    first_y = 9999.0
    first_x = 0.0
    for page_index in sorted(verbatim_pages):
        blocks = extract_verbatim_blocks(page_blocks_cache[page_index])
        for block in blocks:
            first_y = min(first_y, block.y0)
            if first_y == block.y0:
                first_x = block.x0
            quote_text, attribution_text = split_verbatim_block(block)
            if quote_text:
                paragraphs.append(quote_text)
            if attribution_text:
                paragraphs.append(attribution_text)

    if not paragraphs:
        return None
    return Article(0, "VERBATIM", "VERBATIM", "", "", first_page + 1, last_page + 1, paragraphs, first_page, first_y if first_y < 9999 else 0.0, first_x)


def extract_mediafocus_candidates(page: fitz.Page, visible: list[BlockInfo]) -> list[ShortArticleCandidate]:
    title_blocks = [block for block in visible if is_short_title_block(block, page.rect.width, "MEDIAFOCUS")]
    if not title_blocks:
        return []

    sorted_titles = sorted(title_blocks, key=lambda block: (block.x0, block.y0))
    title_spans = [candidate_range(block, page.rect.width) for block in sorted_titles]
    columns = cluster_columns(sorted_titles)

    def column_index(block: BlockInfo) -> int:
        return min(range(len(columns)), key=lambda idx: abs(block.x0 - columns[idx])) if columns else 0

    titles_by_column: dict[int, list[tuple[int, BlockInfo]]] = {}
    for idx, title_block in enumerate(sorted_titles):
        titles_by_column.setdefault(column_index(title_block), []).append((idx, title_block))

    subtitle_parts_by_title: list[list[str]] = [[] for _ in sorted_titles]
    member_blocks_by_title: list[list[BlockInfo]] = [[] for _ in sorted_titles]
    body_pool = [block for block in visible if block not in title_blocks and not is_section_label(block) and not is_byline(block)]
    for block in sorted(body_pool, key=lambda item: (item.x0, item.y0)):
        col = column_index(block)
        column_titles = titles_by_column.get(col, [])
        assigned_idx: int | None = None
        for pos, (idx, title_block) in enumerate(column_titles):
            next_start_y = column_titles[pos + 1][1].y0 if pos + 1 < len(column_titles) else float("inf")
            span = title_spans[idx]
            overlap = min(block.x1, span[1]) - max(block.x0, span[0])
            if overlap <= 10 or block.y0 <= title_block.y1 - 2:
                continue
            if block.y0 >= next_start_y - 2:
                continue
            assigned_idx = idx
            break
        if assigned_idx is None:
            continue
        if is_short_meta_block(block) and block.y0 - sorted_titles[assigned_idx].y1 <= 42:
            subtitle_parts_by_title[assigned_idx].append(normalize_text(block.text))
            continue
        if is_caption(block, page.rect.width):
            continue
        member_blocks_by_title[assigned_idx].append(block)

    candidates: list[ShortArticleCandidate] = []
    for idx, title_block in enumerate(sorted_titles):
        member_blocks = member_blocks_by_title[idx]
        if not member_blocks:
            continue
        title, title_subtitle = split_short_title_parts(title_block)
        subtitle_parts = [part for part in [title_subtitle, *subtitle_parts_by_title[idx]] if part]
        candidates.append(ShortArticleCandidate("MEDIAFOCUS", title, " | ".join(dict.fromkeys(subtitle_parts)), page.number, title_block.y0, member_blocks, "column", title_block.x0))
    return candidates


def extract_short_page_candidates(page: fitz.Page, blocks: list[BlockInfo], occupied_title_pages: set[int], fallback_section: str = "") -> list[ShortArticleCandidate]:
    page_number = page.number + 1
    if page_number < 10 or page_number > 15 or page.number in occupied_title_pages:
        return []
    section = page_section_label(blocks) or fallback_section
    if section == "VERBATIM":
        return []
    if section in {"CONTENTS", "SNAPSHOTS", "EDITORIAL"}:
        return []

    visible: list[BlockInfo] = []
    for block in blocks:
        if is_decorative_block(block) or block.y1 > BOTTOM_MARGIN:
            continue
        if block.y0 < 20 and not is_section_label(block):
            continue
        visible.extend(split_short_candidate_blocks(block))

    if section == "MEDIAFOCUS":
        return extract_mediafocus_candidates(page, visible)

    title_blocks = [block for block in visible if is_short_title_block(block, page.rect.width, section)]
    if not title_blocks:
        return []

    if section == "NETIZENWATCH":
        title_block = max(title_blocks, key=lambda block: (block.max_size, -block.y0, block.width))
        body_blocks = [block for block in visible if block != title_block and not is_section_label(block) and not is_byline(block)]
        return [ShortArticleCandidate(section, format_title(title_block.text), "", page.number, title_block.y0, body_blocks, "single", title_block.x0)] if body_blocks else []

    body_pool = [block for block in visible if block not in title_blocks and not is_section_label(block) and not is_byline(block)]
    sorted_titles = sorted(title_blocks, key=lambda block: (block.y0, block.x0))
    title_spans = [candidate_range(block, page.rect.width) for block in sorted_titles]
    subtitle_parts_by_title: list[list[str]] = [[] for _ in sorted_titles]
    member_blocks_by_title: list[list[BlockInfo]] = [[] for _ in sorted_titles]
    ordered_body_blocks = sorted(body_pool, key=lambda block: (block.y0, block.x0))
    block_assignment: dict[tuple[int, int, int], int] = {}

    for block in ordered_body_blocks:
        best_idx: int | None = None
        best_score = float("-inf")
        for idx, title_block in enumerate(sorted_titles):
            if block.y0 <= title_block.y1 - 2:
                continue
            span = title_spans[idx]
            overlap = min(block.x1, span[1]) - max(block.x0, span[0])
            if overlap <= 10:
                continue
            vertical_gap = block.y0 - title_block.y1
            score = overlap - vertical_gap * 0.2
            if is_short_meta_block(block) and vertical_gap <= 42:
                score += 15
            if score > best_score:
                best_score = score
                best_idx = idx
        if best_idx is None:
            continue
        if is_short_meta_block(block) and block.y0 - sorted_titles[best_idx].y1 <= 42:
            subtitle_parts_by_title[best_idx].append(normalize_text(block.text))
            continue
        if is_caption(block, page.rect.width):
            continue
        block_assignment[(block.block_no, int(block.x0), int(block.y0))] = best_idx

    row_groups: list[list[BlockInfo]] = []
    for block in ordered_body_blocks:
        key = (block.block_no, int(block.x0), int(block.y0))
        if key not in block_assignment:
            continue
        if row_groups and abs(row_groups[-1][0].y0 - block.y0) <= 4:
            row_groups[-1].append(block)
        else:
            row_groups.append([block])
    for group in row_groups:
        if len(group) < 2 or max(block.x0 for block in group) - min(block.x0 for block in group) < 120:
            continue
        counts: dict[int, int] = {}
        for block in group:
            idx = block_assignment[(block.block_no, int(block.x0), int(block.y0))]
            counts[idx] = counts.get(idx, 0) + 1
        target_idx, target_count = max(counts.items(), key=lambda item: item[1])
        if target_count < 2:
            continue
        for block in group:
            block_assignment[(block.block_no, int(block.x0), int(block.y0))] = target_idx

    for block in ordered_body_blocks:
        key = (block.block_no, int(block.x0), int(block.y0))
        idx = block_assignment.get(key)
        if idx is not None:
            member_blocks_by_title[idx].append(block)

    candidates: list[ShortArticleCandidate] = []
    for idx, title_block in enumerate(sorted_titles):
        member_blocks = member_blocks_by_title[idx]
        if not member_blocks:
            continue
        title, title_subtitle = split_short_title_parts(title_block)
        subtitle_parts = [part for part in [title_subtitle, *subtitle_parts_by_title[idx]] if part]
        candidates.append(ShortArticleCandidate(section, title, " | ".join(dict.fromkeys(subtitle_parts)), page.number, title_block.y0, member_blocks, "column", title_block.x0))
    return candidates


def extract_articles(pdf_path: Path) -> list[Article]:
    doc = fitz.open(pdf_path)
    seeds: list[ArticleSeed] = []
    seen_pages: set[int] = set()
    page_blocks_cache: dict[int, list[BlockInfo]] = {}
    for page_index in range(doc.page_count):
        blocks = extract_blocks(doc[page_index])
        page_blocks_cache[page_index] = blocks
        seed = detect_article_seed(doc[page_index], blocks)
        if seed and seed.start_page_index not in seen_pages and seed.title:
            seeds.append(seed)
            seen_pages.add(seed.start_page_index)

    seeds.sort(key=lambda item: item.start_page_index)
    occupied_title_pages = {seed.start_page_index for seed in seeds}
    short_candidates: list[ShortArticleCandidate] = []
    last_front_section = ""
    for page_index in range(doc.page_count):
        page_blocks = page_blocks_cache[page_index]
        page_number = page_index + 1
        section = page_section_label(page_blocks)
        fallback = ""
        if 10 <= page_number <= 15:
            if section and section not in {"EDITORIAL", "CONTENTS", "SNAPSHOTS"}:
                last_front_section = section
            fallback = last_front_section
        else:
            last_front_section = ""
        short_candidates.extend(extract_short_page_candidates(doc[page_index], page_blocks, occupied_title_pages, fallback))

    break_pages = detect_break_pages(doc)
    long_articles: list[Article] = []
    for article_number, seed in enumerate(seeds, start=1):
        end_page_index = doc.page_count - 1
        next_start = doc.page_count
        if article_number < len(seeds):
            next_start = seeds[article_number].start_page_index
            end_page_index = next_start - 1
        interim_breaks = [page for page in break_pages if seed.start_page_index < page < next_start]
        if interim_breaks:
            end_page_index = min(end_page_index, min(interim_breaks) - 1)
        long_articles.append(build_article(doc, seed, end_page_index, article_number))

    verbatim_article = build_verbatim_article(page_blocks_cache)
    short_articles = [build_short_article(candidate, 0) for candidate in short_candidates]
    combined = short_articles + long_articles
    if verbatim_article is not None:
        combined.append(verbatim_article)
    combined.sort(key=lambda article: (article.sort_page_index, article.sort_y, article.sort_x, article.end_pdf_page, article.title))
    for idx, article in enumerate(combined, start=1):
        article.index = idx
    return combined


def save_json(articles: list[Article], path: Path) -> None:
    data = [{"index": a.index, "section": a.section, "title": a.title, "subtitle": a.subtitle, "byline": a.byline, "start_pdf_page": a.start_pdf_page, "end_pdf_page": a.end_pdf_page, "paragraphs": a.paragraphs} for a in articles]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def save_docx(articles: list[Article], path: Path, pdf_name: str) -> None:
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Title"].font.name = "Times New Roman"
    document.styles["Title"].font.size = Pt(18)
    document.styles["Heading 1"].font.name = "Times New Roman"
    document.styles["Heading 1"].font.size = Pt(15)
    document.add_heading("Extracted News Articles", level=0)
    intro = document.add_paragraph()
    intro.add_run("Source PDF: ").bold = True
    intro.add_run(pdf_name)
    for article in articles:
        document.add_page_break()
        document.add_heading(article.title, level=1)
        meta = document.add_paragraph()
        meta.add_run(f"Section: {article.section or 'N/A'}").italic = True
        meta.add_run(f" | PDF pages: {article.page_range}").italic = True
        if article.byline:
            meta.add_run(f" | {article.byline}").italic = True
        if article.subtitle:
            sub = document.add_paragraph()
            sub.add_run(article.subtitle).italic = True
        for paragraph in article.paragraphs:
            document.add_paragraph(paragraph)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract article titles and body text from NewsChina PDF.")
    parser.add_argument("--pdf", required=True, help="Input PDF path")
    parser.add_argument("--json", required=True, help="Output JSON path")
    parser.add_argument("--docx", required=True, help="Output DOCX path")
    args = parser.parse_args()
    pdf_path = Path(args.pdf)
    json_path = Path(args.json)
    docx_path = Path(args.docx)
    articles = extract_articles(pdf_path)
    save_json(articles, json_path)
    save_docx(articles, docx_path, pdf_path.name)
    print(f"Extracted {len(articles)} articles")
    for article in articles:
        print(f"{article.index:02d}. p{article.page_range} | {article.title}")


if __name__ == "__main__":
    main()
